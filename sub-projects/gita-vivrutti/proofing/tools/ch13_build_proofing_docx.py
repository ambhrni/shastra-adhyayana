# -*- coding: utf-8 -*-
import re, json, sys
sys.path.insert(0,'.')
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
exec(open('data13.py',encoding='utf-8').read())
from notes13 import N

DEV='०१२३४५६७८९'
def dn(k): return ''.join(DEV[int(c)] for c in str(k))
FONT='Nirmala UI'

doc=Document()
sec=doc.sections[0]
sec.left_margin=sec.right_margin=Cm(2); sec.top_margin=sec.bottom_margin=Cm(1.8)
st=doc.styles['Normal']; st.font.name=FONT; st.font.size=Pt(12)
rpr=st.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
if rf is None:
    from docx.oxml import OxmlElement
    rf=OxmlElement('w:rFonts'); rpr.append(rf)
for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),FONT)

def setfont(run, size=None, bold=None, color=None, sup=None, italic=None):
    run.font.name=FONT
    r=run._element.get_or_add_rPr(); f=r.find(qn('w:rFonts'))
    if f is None:
        from docx.oxml import OxmlElement
        f=OxmlElement('w:rFonts'); r.insert(0,f)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): f.set(qn(a),FONT)
    if size: run.font.size=Pt(size)
    if bold is not None:
        run.bold=bold
        # complex-script bold too
        from docx.oxml import OxmlElement
        if bold:
            b=OxmlElement('w:bCs'); r.append(b)
    if color: run.font.color.rgb=RGBColor(*color)
    if sup: run.font.superscript=True
    if italic: run.italic=True
    return run

def para(space_after=4, indent=None, align=None):
    p=doc.add_paragraph(); pf=p.paragraph_format
    pf.space_after=Pt(space_after); pf.space_before=Pt(0)
    if indent is not None: pf.left_indent=Cm(indent)
    return p

def add_marked(p, text, size=12):
    """text with **bold** and {^N}; returns list of (text,bold,sup) runs"""
    runs=[]
    for part in re.split(r'(\*\*.+?\*\*|\{\^\d+\})', text):
        if not part: continue
        if part.startswith('**'):
            runs.append((part[2:-2],True,False))
        elif part.startswith('{^'):
            runs.append((dn(part[2:-1]),False,True))
        else:
            runs.append((part,False,False))
    for t,b,s in runs:
        setfont(p.add_run(t), size=size, bold=b if b else None, sup=s)
    return runs

# ---------- front matter
p=para(6); setfont(p.add_run('गीता विवृत्ति — अध्याय १३ (क्षेत्रक्षेत्रज्ञविभागयोगः) — श्लोक १–१२ — संस्कृत पाठ-शोधन पत्रक'), size=16, bold=True)
p=para(2); setfont(p.add_run('Gītā Vivṛtti · Chapter XIII · Verses 13.1–13.12 · Sanskrit proofing sheet (Block A of 3) · Śrī Rāghavendra Tīrtha (रा) only'), size=10, color=(90,90,90))
inst=[
 "HOW TO USE",
 "1. Turn on Track Changes (Review → Track Changes) before you start. Edit the Sanskrit directly; leave the grey [13.N.sK] sentence IDs untouched — they let me match every edit to the right sentence.",
 "2. Bold = the pratīka / lemma bold as I read it in the print. Fix bold with Ctrl+B. Bold in text I did not mark is fine to add.",
 "3. A raised Devanagari numeral (e.g. ²) is a footnote mark in the print; its text is listed under the sentence/verse with the word it attaches to.",
 "4. Under each verse are 'Notes' (in English): where OCR disagreed with the print, where I joined a line-break, and bold I am unsure about. Add Word comments anywhere you want a second look.",
 "5. Only रा's commentary and the मूल verse are included. Vidyādhirāja's (वि) text is excluded; footnotes sitting on वि's text are skipped.",
 "6. Typeface: this print draws ष almost like प, ञ्च like श्च and क्व like क्क. I typed the intended letters. Where the print itself seems to differ from the expected Sanskrit I kept the print and marked the note with ⚑ — please decide those.",
 "7. Each sentence is one daṇḍa-delimited sentence (the unit that becomes one line in the HTML). Verse numbers (॥N॥) are not part of the sentences.",
 "8. When done, save and send the file back; I diff it against my original, apply your changes, then build the English + HTML from your corrected Sanskrit.",
]
for i,t in enumerate(inst):
    p=para(2); setfont(p.add_run(t), size=10, bold=(i==0), color=(60,60,60))
para(6)

sidecar={'chapter':13,'block':'A','verses':{}}
for n in sorted(V):
    d=V[n]
    p=para(4); p.paragraph_format.space_before=Pt(14); p.paragraph_format.keep_with_next=True
    if n==0:
        setfont(p.add_run('अध्यायशीर्षकम्'), size=14, bold=True)
        setfont(p.add_run(f'   (13.0 · chapter title)   {d["pages"]}'), size=9, color=(110,110,110))
    else:
        setfont(p.add_run(f'श्लोक १३.{dn(n)}'), size=14, bold=True)
        setfont(p.add_run(f'   (13.{n})   {d["pages"]}'), size=9, color=(110,110,110))
    # mula
    mula_runs=[]
    for i,l in enumerate(d['mula']):
        p=para(1 if i<len(d['mula'])-1 else 6, indent=1.0); p.paragraph_format.keep_with_next=True
        mula_runs.append(add_marked(p,l,12.5))
        for r in p.runs: r.bold=True
    # ra
    p=para(2); p.paragraph_format.keep_with_next=True
    if d['ra']: setfont(p.add_run('रा-व्याख्या'), size=10, bold=True, color=(120,60,0))
    sidecar['verses'][str(n)]={'pages':d['pages'],'mula':d['mula'],'sentences':[],'footnotes':[{'n':k,'on':on,'text':t} for k,on,t in d['fns']],'notes':N[n]}
    for i,t in enumerate(d['ra'],1):
        sid=f'[13.{n}.s{i}]'
        p=para(3, indent=0.5)
        setfont(p.add_run(sid+' '), size=8, color=(150,150,150))
        add_marked(p,t,12)
        sidecar['verses'][str(n)]['sentences'].append({'id':sid,'text':t})
    if d['fns']:
        p=para(1); setfont(p.add_run('पाद-टिप्पणी'), size=10, bold=True, color=(120,60,0))
        for k,on,t in d['fns']:
            p=para(2, indent=0.5)
            setfont(p.add_run(f'{dn(k)}. '), size=11, bold=True)
            setfont(p.add_run(f'[{on}] '), size=9, color=(110,110,110))
            setfont(p.add_run(t), size=11)
    if N[n]:
        p=para(1); setfont(p.add_run('Notes'), size=9, bold=True, color=(40,90,40))
        for t in N[n]:
            p=para(1, indent=0.5); setfont(p.add_run('▸ '+t), size=9, color=(40,90,40))

out='/mnt/user-data/outputs/ch13_blockA_proofing.docx'
doc.save(out)
json.dump(sidecar, open('/mnt/user-data/outputs/ch13_blockA_proofing_original.json','w',encoding='utf-8'), ensure_ascii=False, indent=1)
print('saved', len(sidecar['verses']), sum(len(v['sentences']) for v in sidecar['verses'].values()))
